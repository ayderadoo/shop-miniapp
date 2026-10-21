"""Browser regression checks with isolated Telegram/API fixtures; no real messages or orders."""
import copy
import os
import json
from pathlib import Path
from playwright.sync_api import sync_playwright, expect

BASE_URL=os.environ.get('RAMP_TEST_URL','http://127.0.0.1:8765')
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'preview'
OUT.mkdir(exist_ok=True)
STATUS = dict(new='Ждёт оплаты',review='Чек на проверке',paid='Оплачен',bought='Собирается',shipping='В пути',arrived='Прибыл',delivered='Доставлен',cancelled='Отменён')
PRODUCTS = [dict(id=1,name='Настольная лампа',category='Дом',price=125000,photo='img/lamp.jpg',description='Тёплый свет для уютных вечеров.',in_stock=True,active=True,options=[],specs=[]),
 dict(id=2,name='Худи оверсайз',category='Одежда',price=89000,photo='img/hoodie.jpg',description='Мягкий хлопок, свободная посадка.',in_stock=True,active=True,options=[dict(name='Размер',values=[dict(label='M',add=0),dict(label='XL',add=10000)])],specs=[dict(name='Материал',value='Хлопок')]),
 dict(id=3,name='Беспроводные наушники',category='Гаджеты',price=210000,photo='img/headphones.jpg',description='Музыка с вами каждый день.',in_stock=True,active=True,options=[],specs=[])]
PROFILE=dict(telegram_id=101,public_id='RAMP-100101',delivery_id='RAMP-D000001',phone_verified=True,subscription_verified=True,user_no=1,first_name='Азиз',last_name='Каримов',username='qa_user',full_name='Каримов Азиз',tg_phone='+998901234567',phone='+998901234567',city='Ташкент',address='Улица Тестовая, дом 1',catalog_view='grid',created_at='2026-10-20T00:00:00Z',blocked=False)
RULES=dict(welcomePct=10,welcomeDays=30,inviterPct=33,inviteeNextPct=11,minQualifyingOrder=200000)
PROMO=dict(current=dict(pct=10,reason='Приветственная скидка'),welcome=dict(active=True,pct=10,until='2026-11-20T00:00:00Z'),coupons=[],invited=0,ref_link='https://t.me/ramp_uzbot?start=invite_00112233445566778899aabbccddeeff',can_enter_ref=True,referred=False,ref_rewarded=False,rules=RULES)
TG="""window.Telegram={WebApp:{initData:'isolated-qa',initDataUnsafe:{user:{id:101,first_name:'Азиз',last_name:'Каримов'}},colorScheme:'light',ready(){},expand(){},onEvent(){},setHeaderColor(){},setBackgroundColor(){},setBottomBarColor(){},BackButton:{show(){},hide(){},onClick(){}},HapticFeedback:{impactOccurred(){},notificationOccurred(){}},showConfirm(t,cb){cb(true)},requestContact(cb){window.qaContact=true;cb(true)}}};"""

class Fixture:
 def __init__(self):
  self.profile=copy.deepcopy(PROFILE);self.cart=[];self.orders=[];self.products=copy.deepcopy(PRODUCTS);self.saved_product=None
  for product in self.products:product['sku']='TEST-'+str(product['id'])
 def route(self,route):
  req=route.request
  if req.method=='OPTIONS':return route.fulfill(status=204,headers={'Access-Control-Allow-Origin':'*','Access-Control-Allow-Headers':'*'})
  if '/rest/v1/products' in req.url:return route.fulfill(json=self.products,headers={'Access-Control-Allow-Origin':'*'})
  body=req.post_data_json or {};action=body.get('action');admin='/functions/v1/admin' in req.url
  if admin:
   data=self.admin(action,body)
  elif action=='me':data=dict(profile=self.profile,cart=self.cart,wish=[],promo=PROMO,payments={'card':'Перевод на карту'},is_admin=True,support='ramphelper')
  elif action=='orders':data={'orders':self.orders}
  elif action=='setPrefs':self.profile['catalog_view']=body['catalog_view'];data={'ok':True}
  elif action=='cartSet':
   product=next(p for p in self.products if p['id']==body['product_id']);opts=body.get('opts',{});variant=', '.join(f"{o['name']}: {opts[o['name']]}" for o in product['options'])
   self.cart=[c for c in self.cart if not(c['product']['id']==product['id'] and c['variant']==variant)]
   if body['qty']>0:self.cart.append(dict(product=product,qty=body['qty'],opts=opts,variant=variant,unit_price=product['price']+sum(next(v for v in o['values'] if v['label']==opts[o['name']]).get('add',0) for o in product['options']),available=True))
   data={'cart':self.cart}
  elif action=='cart':data={'cart':self.cart}
  elif action=='quote':
   total=sum(c['unit_price']*c['qty'] for c in self.cart);data={'quote':dict(subtotal=total,discount_pct=10,discount_amount=round(total*.1),discount_reason='Приветственная скидка',total=round(total*.9))}
  elif action=='checkout':
   assert body.get('request_id'),'Missing checkout idempotency key'
   total=sum(c['unit_price']*c['qty'] for c in self.cart)
   self.orders=[dict(id=1,items=[dict(name=c['product']['name'],price=c['unit_price'],qty=c['qty'],variant=c['variant']) for c in self.cart],total=round(total*.9),subtotal=total,discount_pct=10,discount_amount=round(total*.1),discount_reason='Приветственная скидка',status='new',payment_method='card',pay_card='8600 0000 0000 0001',created_at='2026-10-20T00:00:00Z')];self.cart=[]
   data=dict(order=self.orders[0],orders=self.orders,profile=self.profile,promo=PROMO,cart=[])
  elif action=='photoRequests':data={'requests':[]}
  else:data={'ok':True}
  route.fulfill(json=data,headers={'Access-Control-Allow-Origin':'*'})
 def admin(self,action,b):
  order=dict(id=12,telegram_id=101,status='review',created_at='2026-10-20T10:00:00Z',full_name='Каримов Азиз',phone='+998901234567',address='Ташкент',items=[dict(name='Худи',qty=1,price=89000,variant='Размер: M')],total=80100,subtotal=89000,discount_pct=10,discount_amount=8900,payment_method='card',user=self.profile)
  if action=='stats':return dict(byStatus={'review':1,'new':2},users=3,newRequests=1,revenue30=540000,statuses=STATUS)
  if action=='orders':return dict(orders=[order],payments={'card':'Карта'},has_more=False)
  if action=='products':return {'products':self.products}
  if action=='users':return dict(users=[self.profile],has_more=False)
  if action=='user':return dict(user=self.profile,orders=[order],coupons=[],invited=[],referrer=None,spent=0,isAdmin=False)
  if action=='requests':return dict(requests=[],has_more=False)
  if action=='settings':return dict(promo=RULES,payments={'card':'Карта'},cards=['8600 0000 0000 0001'],admins=[dict(telegram_id=101,name='Админ')],me=101,bots=dict(shop=dict(configured=False,username='ramp_uzbot'),admin=dict(configured=True,username='qa_adminbot',webhook_ok=True)))
  if action=='exportData':return dict(rows=([self.profile] if b['entity']=='users' else self.products if b['entity']=='products' else [order] if b['entity']=='orders' else []),next_offset=None)
  if action=='saveProduct':self.saved_product=b;return dict(product=b)
  return {'ok':True}

def main():
 with sync_playwright() as p:
  browser=p.chromium.launch(headless=True,args=['--no-sandbox'])
  context=browser.new_context(viewport={'width':390,'height':844},device_scale_factor=1)
  fixture=Fixture();context.route('https://telegram.org/**',lambda route:route.fulfill(body=TG,content_type='application/javascript'))
  context.route('https://*.supabase.co/**',fixture.route)
  page=context.new_page();errors=[];page.on('pageerror',lambda error:errors.append(str(error)))
  page.goto(BASE_URL,wait_until='networkidle');expect(page.locator('#grid .card')).to_have_count(3)
  expect(page.locator('#gate')).not_to_have_class('gate show')
  page.locator('#viewToggle [data-view="list"]').click();expect(page.locator('#grid')).to_have_class('grid list')
  page.locator('#viewToggle [data-view="grid"]').click()
  page.screenshot(path=str(OUT/'shop-catalog.png'),full_page=True,animations="disabled")
  page.locator('#grid [data-add="2"]').click();expect(page.locator('#psAction button')).to_be_disabled()
  page.locator('[data-option="0"][data-value="1"]').click();expect(page.locator('#psPrice')).to_have_text('99 000 сум')
  page.locator('#psAction [data-act="add"]').click();expect(page.locator('#cartCount')).to_have_text('1')
  page.locator('#psAction [data-act="inc"]').click();page.locator('#psAction [data-act="inc"]').click();expect(page.locator('#cartCount')).to_have_text('3')
  page.locator('#psAction [data-act="toCart"]').click();expect(page.locator('#cartBox')).to_contain_text('Размер: XL')
  page.locator('#toCheckout').click();expect(page.locator('#coSubmit')).to_be_enabled();page.locator('#coSubmit').click()
  expect(page.locator('#ordersBox')).to_contain_text('Заказ №1')
  assert len(fixture.orders)==1 and fixture.orders[0]['items'][0]['qty']==3
  page.locator('.tab[data-tab="profile"]').click();expect(page.locator('#deliveryId')).to_have_text('RAMP-D000001');page.screenshot(path=str(OUT/'shop-profile.png'),full_page=True,animations="disabled")
  page.set_viewport_size({'width':320,'height':740});assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'),'Shop horizontal overflow'
  page.goto(BASE_URL+'/admin.html',wait_until='networkidle');expect(page.locator('#oList .card')).to_have_count(1)
  page.set_viewport_size({'width':1280,'height':850});page.screenshot(path=str(OUT/'admin-desktop.png'),full_page=True,animations="disabled")
  page.locator('[data-tab="users"]').click();page.locator('#uList [data-user]').click();page.locator('[data-uorder]').click();expect(page.locator('#shTitle')).to_have_text('Заказ №12')
  page.locator('[data-close-sheet]').click();expect(page.locator('#shTitle')).to_contain_text('Каримов');page.locator('[data-close-sheet]').click()
  page.locator('[data-tab="products"]').click();page.locator('[data-product="2"]').click();expect(page.locator('[data-option-name]')).to_have_value('Размер')
  page.locator('#epName').fill('Худи — новая версия');page.locator('#shBar [data-act="saveProduct"]').click();expect(page.locator('#sheet')).not_to_have_class('sheet open')
  assert fixture.saved_product['options'][0]['values'][1]['add']==10000,'Variant lost on admin save'
  page.locator('[data-tab="settings"]').click();expect(page.locator('#sBox')).to_contain_text('@ramp_uzbot');expect(page.locator('#shopBotToken')).to_have_attribute('type','password')
  with page.expect_download() as d:page.locator('#exportAll').click()
  d.value.save_as(str(OUT/'export-check.xlsx'))
  from openpyxl import load_workbook
  book=load_workbook(OUT/'export-check.xlsx');assert book.sheetnames==['Заказы','Клиенты','Товары','Заявки','Купоны','Реестр ID','Закупка','Истории','Отзывы'];headers=[cell.value for cell in book['Клиенты'][1]];assert book['Клиенты'].cell(2,headers.index('Личный ID')+1).value=='RAMP-D000001'
  with page.expect_download() as d:page.locator('[data-act="exportCsv"]').click()
  d.value.save_as(str(OUT/'export-check.csv'));assert (OUT/'export-check.csv').read_bytes().startswith(b'\xef\xbb\xbf')
  page.set_viewport_size({'width':390,'height':844});page.locator('[data-tab="products"]').click();page.screenshot(path=str(OUT/'admin-mobile.png'),full_page=True,animations="disabled")
  page.set_viewport_size({'width':320,'height':740});assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'),'Admin horizontal overflow'
  context.close();browser.close()
  assert not errors,errors
  print('PASS: shop variants, rapid cart updates, checkout, profile ID, responsive layouts, admin back navigation, product editor, XLSX and CSV')

if __name__=='__main__':main()
